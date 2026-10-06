import io
import openpyxl
from werkzeug.security import generate_password_hash
from models import db, Member
from services.qr_service import generate_pin, generate_qr_code
from services.email_service import log_action

def generate_members_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Membres C-TECH"

    headers = ["ID", "N° Membre", "Nom", "Prénom", "Filière", "Niveau", "Email", "Téléphone", "Statut", "Bureau", "Date Inscription"]
    ws.append(headers)

    members = Member.query.order_by(Member.created_at.desc()).all()
    for m in members:
        ws.append([
            m.id,
            m.member_number,
            m.last_name,
            m.first_name,
            m.major,
            m.level,
            m.email,
            m.phone,
            m.status,
            "Oui" if m.is_bureau else "Non",
            m.created_at.strftime("%Y-%m-%d %H:%M") if m.created_at else ""
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def import_members_from_excel(file_storage, operator="Bureau"):
    wb = openpyxl.load_workbook(file_storage)
    ws = wb.active

    imported = 0
    errors = []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0, ["Le fichier Excel est vide."]

    # Skip header if present
    start_idx = 1 if isinstance(rows[0][0], str) and "nom" in str(rows[0][0]).lower() or "id" in str(rows[0][0]).lower() else 0

    from datetime import datetime, timedelta
    from blueprints.auth import generate_member_number

    for idx, row in enumerate(rows[start_idx:], start=start_idx + 1):
        try:
            if not row or len(row) < 5 or not row[0]:
                continue
            
            # Simple layout detection: Nom, Prénom, Filière, Niveau, Email, Téléphone
            last_name = str(row[0]).strip()
            first_name = str(row[1]).strip() if len(row) > 1 and row[1] else "Étudiant"
            major = str(row[2]).strip() if len(row) > 2 and row[2] else "Informatique"
            level = str(row[3]).strip() if len(row) > 3 and row[3] else "Licence"
            email = str(row[4]).strip() if len(row) > 4 and row[4] else f"user{idx}@clubtech.org"
            phone = str(row[5]).strip() if len(row) > 5 and row[5] else "0000000000"

            if Member.query.filter_by(email=email).first():
                errors.append(f"Ligne {idx}: Email {email} déjà existant.")
                continue

            pin_code = generate_pin()
            member_no = generate_member_number()
            qr_rel = generate_qr_code(member_no)
            expires_at = datetime.utcnow() + timedelta(hours=24)

            new_m = Member(
                member_number=member_no,
                pin=generate_password_hash(pin_code),
                last_name=last_name,
                first_name=first_name,
                major=major,
                level=level,
                email=email,
                phone=phone,
                qr_code_path=qr_rel,
                must_change_pin=True,
                pin_expires_at=expires_at
            )
            db.session.add(new_m)
            imported += 1
        except Exception as e:
            errors.append(f"Ligne {idx}: {str(e)}")

    db.session.commit()
    log_action(operator, f"Importation de {imported} membres depuis Excel.")
    return imported, errors
